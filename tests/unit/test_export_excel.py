"""Teste l'export Excel réellement (ouverture avec openpyxl). Aucun appel
LLM, entièrement déterministe."""

from __future__ import annotations

import openpyxl
import pandas as pd

from agent.export.excel import write_excel


def _sample_df() -> pd.DataFrame:
    return pd.DataFrame({"category": ["Beverages", "Dairy"], "revenue": [12000.5, 15300.0]})


def test_write_excel_creates_valid_file(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = write_excel(tmp_path / "report.xlsx", sheets={"Donnees": _sample_df()})
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert "Donnees" in wb.sheetnames


def test_write_excel_writes_header_and_data(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = write_excel(tmp_path / "report.xlsx", sheets={"Donnees": _sample_df()})
    wb = openpyxl.load_workbook(path)
    ws = wb["Donnees"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("category", "revenue")
    assert rows[1] == ("Beverages", 12000.5)
    assert rows[2] == ("Dairy", 15300.0)


def test_write_excel_multiple_sheets(tmp_path) -> None:  # type: ignore[no-untyped-def]
    other_df = pd.DataFrame({"order_id": [1, 2], "total": [100, 200]})
    path = write_excel(
        tmp_path / "report.xlsx", sheets={"Donnees": _sample_df(), "Commandes": other_df}
    )
    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"Donnees", "Commandes"}


def test_write_excel_adds_native_chart(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = write_excel(
        tmp_path / "report.xlsx",
        sheets={"Donnees": _sample_df()},
        chart_sheet="Donnees",
        chart_title="Revenu par categorie",
    )
    wb = openpyxl.load_workbook(path)
    ws = wb["Donnees"]
    assert len(ws._charts) == 1  # noqa: SLF001 - seule façon d'inspecter les graphiques via openpyxl


def test_write_excel_adds_metadata_sheet(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = write_excel(
        tmp_path / "report.xlsx",
        sheets={"Donnees": _sample_df()},
        metadata={"question": "CA par categorie ?", "sql": "SELECT ..."},
    )
    wb = openpyxl.load_workbook(path)
    assert "Metadonnees" in wb.sheetnames
    ws = wb["Metadonnees"]
    rows = list(ws.iter_rows(values_only=True))
    assert ("question", "CA par categorie ?") in rows


def test_write_excel_sanitizes_invalid_sheet_names(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = write_excel(tmp_path / "report.xlsx", sheets={"Ventes/Region:2026": _sample_df()})
    wb = openpyxl.load_workbook(path)
    assert all("/" not in name and ":" not in name for name in wb.sheetnames)


def test_write_excel_handles_empty_dataframe(tmp_path) -> None:  # type: ignore[no-untyped-def]
    empty_df = pd.DataFrame({"category": [], "revenue": []})
    path = write_excel(
        tmp_path / "report.xlsx", sheets={"Donnees": empty_df}, chart_sheet="Donnees"
    )
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert "Donnees" in wb.sheetnames
